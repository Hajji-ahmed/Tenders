"""Export d'un document de candidature en DOCX (Phase 9).

Ce fichier est ce que l'acheteur recevra : page de garde, une section par titre, et le Markdown
rendu en paragraphes, listes et tableaux. La mention « assistance IA » n'apparaît qu'une fois le
document validé par un humain (RB-006) — un brouillon porte au contraire un bandeau qui dit qu'il
n'a pas encore été relu, pour qu'un envoi par mégarde se voie."""

import io
import re
from datetime import date

from docx import Document as new_document
from docx.document import Document  # le type ; `docx.Document` est la fabrique
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from app.models import AppDocStatus, ApplicationDocument
from app.services.company_facts import CompanyFacts

AI_NOTICE = "Document généré avec assistance IA"
DRAFT_NOTICE = "Brouillon — non validé : à relire avant envoi"
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_HEADING = re.compile(r"^(#{2,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _write_rich(paragraph: Paragraph, text: str) -> None:
    """Rend le gras Markdown (`**…**`) ; le reste du balisage est écrit tel quel, sans le perdre."""
    for index, part in enumerate(_BOLD.split(text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1


def _add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for cell_index, value in enumerate(row):
            cell = table.cell(row_index, cell_index)
            cell.text = ""
            _write_rich(cell.paragraphs[0], value)
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def write_markdown(document: Document, markdown: str) -> None:
    """Markdown simple → Word : sous-titres, listes, tableaux, paragraphes. Ce qui n'est pas reconnu
    reste du texte : mieux vaut une ligne littérale qu'une ligne disparue."""
    table: list[list[str]] = []
    for line in (markdown or "").splitlines():
        if _TABLE_ROW.match(line):
            if not _TABLE_SEPARATOR.match(line):
                table.append(_cells(line))
            continue
        if table:
            _add_table(document, table)
            table = []
        stripped = line.strip()
        if not stripped:
            continue
        if heading := _HEADING.match(stripped):
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
        elif bullet := _BULLET.match(stripped):
            _write_rich(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered := _NUMBERED.match(stripped):
            _write_rich(document.add_paragraph(style="List Number"), numbered.group(1))
        else:
            _write_rich(document.add_paragraph(), stripped)
    if table:
        _add_table(document, table)


def _cover(document: Document, doc: ApplicationDocument, company: CompanyFacts, tender_title: str) -> None:
    title = document.add_heading(doc.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in (
        company.trade_name or company.legal_name or "",
        tender_title,
        f"Version {doc.current_version} — {date.today():%d/%m/%Y}",
    ):
        if not text:
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size = Pt(12)
    if doc.status != AppDocStatus.validated:
        warning = document.add_paragraph()
        warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = warning.add_run(DRAFT_NOTICE)
        run.bold = True
    document.add_page_break()


def _footer(document: Document, doc: ApplicationDocument) -> None:
    """RB-006 : la mention d'assistance IA accompagne un document validé par un humain."""
    if doc.status != AppDocStatus.validated:
        return
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    validated = doc.updated_at.date() if doc.updated_at else date.today()
    paragraph.text = f"{AI_NOTICE} — validé le {validated:%d/%m/%Y}"


def export_docx(doc: ApplicationDocument, company: CompanyFacts, *, tender_title: str = "") -> bytes:
    """Le document complet, prêt à être déposé : page de garde, une section par titre."""
    document = new_document()
    _cover(document, doc, company, tender_title)
    for section in doc.sections:
        document.add_heading(section.title, level=1)
        write_markdown(document, section.content_md or "")
    _footer(document, doc)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_key(doc: ApplicationDocument) -> str:
    return f"applications/{doc.application_id}/{doc.id}-v{doc.current_version}.docx"
