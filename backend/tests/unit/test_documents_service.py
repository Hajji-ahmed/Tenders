"""Service documents (Tâche 2.3) : validation des envois, stockage, versions, archivage, RB-007."""

import hashlib
import io
import zipfile
from datetime import date, timedelta

import openpyxl
import pymupdf
import pytest
from docx import Document
from sqlalchemy import func, select

from app.core.config import get_settings
from app.core.errors import ConflictError, ValidationError
from app.models.audit import AuditLog
from app.models.document import CompanyDocument, DocumentCategory, DocumentStatus, ExtractionStatus
from app.services.documents import ALLOWED_MIMES, DocumentService, refresh_expiry_statuses, validate_upload

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TXT = "text/plain"
ZIP = "application/zip"

YESTERDAY = date.today() - timedelta(days=1)
TOMORROW = date.today() + timedelta(days=1)


def _audit(db, action: str) -> list[AuditLog]:
    return list(db.scalars(select(AuditLog).where(AuditLog.action == action)))


def _count_documents(db) -> int:
    return db.scalar(select(func.count()).select_from(CompanyDocument))


# --- validate_upload -----------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "mime"),
    [
        ("sample.pdf", PDF),
        ("sample.docx", DOCX),
        ("sample.xlsx", XLSX),
        ("sample.txt", TXT),
        ("sample.zip", ZIP),
        ("sample.zip", "application/x-zip-compressed"),
    ],
)
def test_validate_upload_accepts_allowed_formats(fixtures_dir, name, mime):
    data = (fixtures_dir / name).read_bytes()
    assert validate_upload(name, data, mime) == hashlib.sha256(data).hexdigest()


def test_allowed_mimes_cover_spec_formats():
    assert set(ALLOWED_MIMES.values()) == {".pdf", ".docx", ".xlsx", ".txt", ".zip"}


@pytest.mark.parametrize(
    ("filename", "data", "mime"),
    [
        ("virus.exe", b"MZ\x90\x00", "application/octet-stream"),  # type non autorisé
        ("photo.png", b"\x89PNG\r\n\x1a\n", "image/png"),  # type non autorisé
        ("cv.docx", b"%PDF-1.4", PDF),  # extension incohérente avec le MIME
        ("kbis", b"%PDF-1.4", PDF),  # pas d'extension
        ("piege.pdf", b"MZ\x90\x00" + b"\x00" * 64, PDF),  # exécutable renommé en .pdf
        ("vide.txt", b"", TXT),  # fichier vide
    ],
)
def test_validate_upload_rejects_forbidden_input(filename, data, mime):
    with pytest.raises(ValidationError) as exc:
        validate_upload(filename, data, mime)
    assert exc.value.status_code == 422


def test_validate_upload_rejects_oversize(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    limit = 1024 * 1024
    assert validate_upload("ok.txt", b"a" * limit, TXT)  # exactement 1 Mo : accepté
    with pytest.raises(ValidationError) as exc:
        validate_upload("trop.txt", b"a" * (limit + 1), TXT)
    assert exc.value.code == "file_too_large"
    assert "1 Mo" in exc.value.message


# --- upload / new_version / versions -------------------------------------------------------------


def test_upload_and_version(db, company, storage, user, fixtures_dir):
    svc = DocumentService(db, storage, user_id=user.id, author=user.email)
    v1 = (fixtures_dir / "sample.pdf").read_bytes()
    doc = svc.upload(
        filename="attestation-fiscale.pdf",
        data=v1,
        content_type=PDF,
        category=DocumentCategory.attestation,
        name="Attestation fiscale 2026",
        tags=["fiscal", "2026"],
        expires_at=date(2026, 12, 31),
    )
    assert doc.company_id == company.id
    assert doc.version == 1
    assert doc.name == "Attestation fiscale 2026"
    assert doc.status == DocumentStatus.valid
    assert doc.extraction_status == ExtractionStatus.pending
    assert doc.sha256 == hashlib.sha256(v1).hexdigest()
    assert doc.storage_key == f"company/{doc.sha256[:2]}/{doc.sha256}.pdf"
    assert doc.size_bytes == len(v1)
    assert doc.tags == ["fiscal", "2026"]
    assert doc.is_usable
    assert storage.exists(doc.storage_key) and storage.get(doc.storage_key) == v1
    versions = svc.versions(doc)
    assert [v.version_number for v in versions] == [1]
    assert versions[0].author == user.email  # libellé lisible dans l'historique, pas l'UUID
    assert len(_audit(db, "document.uploaded")) == 1

    v2 = b"%PDF-1.4 nouvelle version"
    key_v1 = doc.storage_key
    doc.extracted_text = "texte v1"
    doc = svc.new_version(doc, filename="attestation-2026.pdf", data=v2, content_type=PDF, changelog="MAJ")
    assert doc.version == 2
    assert doc.sha256 == hashlib.sha256(v2).hexdigest()
    assert doc.storage_key != key_v1
    assert doc.extraction_status == ExtractionStatus.pending and doc.extracted_text is None
    assert storage.get(doc.storage_key) == v2 and storage.exists(key_v1)  # l'ancienne version reste
    versions = svc.versions(doc)
    assert [v.version_number for v in versions] == [1, 2]
    assert versions[0].storage_key == key_v1 and versions[1].storage_key == doc.storage_key
    assert versions[1].changelog == "MAJ"
    assert len(_audit(db, "document.versioned")) == 1
    assert _count_documents(db) == 1


def test_upload_rejects_duplicate_same_category(db, company, storage):
    svc = DocumentService(db, storage)
    same = b"%PDF-1.4 same"
    first = svc.upload(filename="a.pdf", data=same, content_type=PDF, category=DocumentCategory.autre)
    with pytest.raises(ConflictError) as exc:
        svc.upload(filename="b.pdf", data=same, content_type=PDF, category=DocumentCategory.autre)
    assert exc.value.status_code == 409 and "a.pdf" in exc.value.message
    # Même contenu dans une autre catégorie : autorisé (ex. une attestation aussi classée « référence »).
    other = svc.upload(filename="b.pdf", data=same, content_type=PDF, category=DocumentCategory.reference)
    assert other.storage_key == first.storage_key  # clé déterministe : un seul objet stocké
    # Une fois l'original archivé, le même fichier peut être ré-uploadé dans sa catégorie.
    svc.archive(first)
    again = svc.upload(filename="a.pdf", data=same, content_type=PDF, category=DocumentCategory.autre)
    assert again.id != first.id
    assert _count_documents(db) == 3


def test_upload_rejects_bad_type_without_side_effects(db, company, storage):
    svc = DocumentService(db, storage)
    with pytest.raises(ValidationError):
        svc.upload(filename="x.exe", data=b"MZ", content_type="application/octet-stream", category="autre")
    assert _count_documents(db) == 0
    assert _audit(db, "document.uploaded") == []


def test_upload_ignores_none_metadata(db, company, storage):
    """L'endpoint transmet des champs de formulaire optionnels : None ne doit pas écraser les défauts."""
    svc = DocumentService(db, storage)
    doc = svc.upload(
        filename="note.txt",
        data=b"note",
        content_type=TXT,
        category=DocumentCategory.autre,
        name=None,
        description=None,
        issued_at=None,
        expires_at=None,
        tags=None,
    )
    assert doc.name == "note.txt" and doc.tags == [] and doc.description is None


def test_upload_rejects_unknown_metadata(db, company, storage):
    svc = DocumentService(db, storage)
    with pytest.raises(TypeError):
        svc.upload(filename="a.txt", data=b"a", content_type=TXT, category=DocumentCategory.autre, owner="x")


# --- archive -------------------------------------------------------------------------------------


def test_archive(db, company, storage, user):
    svc = DocumentService(db, storage, user_id=user.id)
    doc = svc.upload(filename="kbis.pdf", data=b"%PDF-1.4 kbis", content_type=PDF, category="administratif")
    svc.archive(doc)
    assert doc.status == DocumentStatus.archived and not doc.is_usable
    assert storage.exists(doc.storage_key)  # archivage logique : le fichier est conservé
    svc.archive(doc)  # idempotent
    assert len(_audit(db, "document.archived")) == 1
    with pytest.raises(ConflictError):
        svc.new_version(doc, filename="kbis.pdf", data=b"%PDF-1.4 v2", content_type=PDF)


# --- RB-007 : expiration -------------------------------------------------------------------------


def test_is_usable_false_on_expired_document(db, company, storage):
    svc = DocumentService(db, storage)
    doc = svc.upload(
        filename="att.pdf",
        data=b"%PDF-1.4 att",
        content_type=PDF,
        category="attestation",
        expires_at=YESTERDAY,
    )
    assert doc.status == DocumentStatus.valid  # le statut n'est recalculé que par la tâche planifiée…
    assert doc.is_expired and not doc.is_usable  # …mais RB-007 s'applique immédiatement via is_usable
    ok = svc.upload(
        filename="ok.pdf", data=b"%PDF-1.4 ok", content_type=PDF, category="attestation", expires_at=TOMORROW
    )
    assert not ok.is_expired and ok.is_usable


def test_refresh_expiry_statuses(db, company, storage):
    svc = DocumentService(db, storage)
    up = DocumentCategory.attestation
    expired = svc.upload(
        filename="a.pdf", data=b"%PDF a", content_type=PDF, category=up, expires_at=YESTERDAY
    )
    future = svc.upload(filename="b.pdf", data=b"%PDF b", content_type=PDF, category=up, expires_at=TOMORROW)
    no_expiry = svc.upload(filename="c.pdf", data=b"%PDF c", content_type=PDF, category=up)
    draft = svc.upload(
        filename="d.pdf", data=b"%PDF d", content_type=PDF, category=up, expires_at=YESTERDAY, status="draft"
    )
    archived = svc.upload(
        filename="e.pdf", data=b"%PDF e", content_type=PDF, category=up, expires_at=YESTERDAY
    )
    svc.archive(archived)

    assert refresh_expiry_statuses(db) == 1
    assert expired.status == DocumentStatus.expired and not expired.is_usable
    assert future.status == DocumentStatus.valid and future.is_usable
    assert no_expiry.status == DocumentStatus.valid and no_expiry.is_usable
    assert draft.status == DocumentStatus.draft
    assert archived.status == DocumentStatus.archived
    assert refresh_expiry_statuses(db) == 0  # idempotent

    # Date prolongée (PATCH) : la prochaine passe remet le document en `valid`.
    expired.expires_at = date.today() + timedelta(days=90)
    db.flush()
    assert refresh_expiry_statuses(db) == 1
    assert expired.status == DocumentStatus.valid and expired.is_usable


# --- fixtures binaires ---------------------------------------------------------------------------


def test_sample_fixtures_are_readable(fixtures_dir):
    """Les fichiers d'exemple sont de vrais documents (réutilisés en Phase 6 pour l'extraction)."""
    with pymupdf.open(stream=(fixtures_dir / "sample.pdf").read_bytes(), filetype="pdf") as pdf:
        assert pdf.page_count == 1
        assert "Attestation fiscale InnoSustain 2026" in pdf[0].get_text()
    paragraphs = [
        p.text for p in Document(io.BytesIO((fixtures_dir / "sample.docx").read_bytes())).paragraphs
    ]
    assert any("InnoSustain" in p for p in paragraphs)
    wb = openpyxl.load_workbook(io.BytesIO((fixtures_dir / "sample.xlsx").read_bytes()))
    assert wb.active["A2"].value == "Audit énergétique"
    assert (fixtures_dir / "sample.txt").read_bytes().startswith(b"Note interne InnoSustain")
    with zipfile.ZipFile(io.BytesIO((fixtures_dir / "sample.zip").read_bytes())) as zf:
        assert sorted(zf.namelist()) == ["sample.pdf", "sample.txt"]
