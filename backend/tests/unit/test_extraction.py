"""Extraction de texte multi-formats : PDF (pages), DOCX (paragraphes + tableaux), XLSX (une page
par feuille), TXT (utf-8 puis latin-1), ZIP (plusieurs documents, un niveau de récursion, entrées
suspectes refusées). Les cas limites sont générés en mémoire ; les fixtures versionnées sont relues."""

import io
import zipfile

import openpyxl
import pymupdf
import pytest
from docx import Document

from app.services.extraction import ExtractionError, extract

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _pdf(pages: list[str], *, password: str | None = None) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 90), text, fontsize=11)
    kwargs = {}
    if password:
        kwargs = dict(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password)
    data = doc.tobytes(**kwargs)
    doc.close()
    return data


def _docx() -> bytes:
    d = Document()
    d.add_paragraph("Règlement de consultation — article 1")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Critère", "Poids"
    table.cell(1, 0).text, table.cell(1, 1).text = "Prix", "40 %"
    d.add_paragraph("Article 2 : remise des offres avant le 20/11/2026.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx(rows_second_sheet: int = 3) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bordereau"
    ws.append(["Poste", "Quantité", "Prix"])
    ws.append(["Luminaire LED", 4500, 1200])
    ws.append([None, None, None])  # ligne vide ignorée
    ws2 = wb.create_sheet("Planning")
    for i in range(rows_second_sheet):
        ws2.append([f"Tâche {i}", i])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def test_pdf_gives_one_entry_per_page_with_text():
    (doc,) = extract(_pdf(["Page une : objet du marché", "Page deux : critères"]), PDF, "dce.pdf")
    assert doc.filename == "dce.pdf" and doc.mime_type == PDF
    assert [p.number for p in doc.pages] == [1, 2]
    assert "objet du marché" in doc.pages[0].text and "critères" in doc.pages[1].text
    assert doc.char_count == sum(len(p.text) for p in doc.pages) > 0


def test_protected_and_scanned_pdfs_raise_readable_errors():
    with pytest.raises(ExtractionError, match="PDF protégé ou illisible"):
        extract(_pdf(["secret"], password="x"), PDF, "protected.pdf")
    with pytest.raises(ExtractionError, match="scan"):
        extract(_pdf(["", ""]), PDF, "scan.pdf")
    with pytest.raises(ExtractionError, match="PDF protégé ou illisible"):
        extract(b"%PDF-1.4 pas un vrai pdf", PDF, "corrupt.pdf")


def test_docx_keeps_paragraphs_and_tables_in_order():
    (doc,) = extract(_docx(), DOCX, "RC.docx")
    assert len(doc.pages) == 1 and doc.pages[0].number == 1
    text = doc.pages[0].text
    assert "article 1" in text and "Critère | Poids" in text and "Prix | 40 %" in text
    assert text.index("article 1") < text.index("Critère | Poids") < text.index("Article 2")


def test_xlsx_gives_one_page_per_sheet_with_tab_separated_rows():
    (doc,) = extract(_xlsx(), XLSX, "bordereau.xlsx")
    assert [p.number for p in doc.pages] == [1, 2]
    assert doc.pages[0].text.splitlines() == ["Poste\tQuantité\tPrix", "Luminaire LED\t4500\t1200"]
    assert doc.pages[1].text.startswith("Tâche 0\t0")
    assert doc.metadata["sheets"] == ["Bordereau", "Planning"]

    (big,) = extract(_xlsx(rows_second_sheet=2500), XLSX, "big.xlsx")
    assert len(big.pages[1].text.splitlines()) == 2000 and big.metadata["truncated_sheets"] == ["Planning"]


def test_txt_decodes_utf8_then_latin1():
    (utf8,) = extract("Délai : 30 jours".encode(), "text/plain", "note.txt")
    (latin,) = extract("Délai : 30 jours".encode("latin-1"), "text/plain", "note.txt")
    assert utf8.pages[0].text == latin.pages[0].text == "Délai : 30 jours"


def test_zip_extracts_supported_members_one_level_deep_and_reports_the_rest():
    inner = _zip({"annexe.txt": b"annexe", "trop-profond.zip": _zip({"x.txt": b"x"})})
    data = _zip(
        {
            "dossier/CCTP.pdf": _pdf(["cahier des charges"]),
            "dossier/RC.docx": _docx(),
            "bordereau.xlsx": _xlsx(),
            "lisez-moi.txt": b"bonjour",
            "installeur.exe": b"MZ\x90",
            "pieces.zip": inner,
        }
    )
    docs = extract(data, "application/zip", "dce.zip")
    assert [d.filename for d in docs] == [
        "dossier/CCTP.pdf",
        "dossier/RC.docx",
        "bordereau.xlsx",
        "lisez-moi.txt",
        "pieces.zip/annexe.txt",
    ]
    assert docs[0].mime_type == PDF and docs[4].metadata["inner_filename"] == "pieces.zip/annexe.txt"
    assert docs[4].metadata["skipped"] == ["pieces.zip/trop-profond.zip"]  # récursion limitée à un niveau
    assert docs[0].metadata["skipped"] == ["installeur.exe"]


def test_zip_refuses_path_traversal_and_reports_unreadable_members():
    with pytest.raises(ExtractionError, match=r"\.\."):
        extract(_zip({"../evil.txt": b"x"}), "application/zip", "evil.zip")
    docs = extract(_zip({"ok.txt": b"ok", "bad.pdf": b"pas un pdf"}), "application/zip", "mixte.zip")
    assert [d.filename for d in docs] == ["ok.txt"]
    assert docs[0].metadata["failed"] == {"bad.pdf": "PDF protégé ou illisible"}
    with pytest.raises(ExtractionError, match="Aucun document exploitable"):
        extract(_zip({"only.exe": b"MZ"}), "application/zip", "vide.zip")


def test_unsupported_type_and_corrupt_archive():
    with pytest.raises(ExtractionError, match="non pris en charge"):
        extract(b"GIF89a", "image/gif", "logo.gif")
    with pytest.raises(ExtractionError, match="Archive illisible"):
        extract(b"PK\x03\x04 broken", "application/zip", "broken.zip")


def test_versioned_fixtures_are_readable(fixtures_dir):
    (pdf,) = extract((fixtures_dir / "sample.pdf").read_bytes(), PDF, "sample.pdf")
    assert len(pdf.pages) == 1 and "Attestation fiscale InnoSustain 2026" in pdf.pages[0].text
    (docx,) = extract((fixtures_dir / "sample.docx").read_bytes(), DOCX, "sample.docx")
    assert "Présentation InnoSustain" in docx.pages[0].text
    (xlsx,) = extract((fixtures_dir / "sample.xlsx").read_bytes(), XLSX, "sample.xlsx")
    assert xlsx.pages[0].text.splitlines()[1] == "Audit énergétique\tOffice National X\tÉnergie\t2025"
    zipped = extract((fixtures_dir / "sample.zip").read_bytes(), "application/zip", "sample.zip")
    assert sorted(d.filename for d in zipped) == ["sample.pdf", "sample.txt"]
