"""Export DOCX (Task 9.4) : le fichier que l'acheteur recevra — page de garde, une section par
titre, Markdown rendu. La mention d'assistance IA n'accompagne qu'un document validé (RB-006) ;
un brouillon porte au contraire un avertissement, pour qu'un envoi par mégarde se voie."""

import io

import pytest
from docx import Document as read_docx

from app.models import (
    AppDocStatus,
    Application,
    ApplicationDocument,
    ApplicationSection,
    DocumentType,
    Tender,
)
from app.services.company_facts import from_company
from app.services.docx_export import AI_NOTICE, DRAFT_NOTICE, export_docx, export_key, write_markdown

MARKDOWN = """### Notre approche

Nous proposons une **rénovation complète** du parc.

- Audit énergétique préalable
- Remplacement des luminaires
1. Première phase
2. Seconde phase

| Phase | Durée | Livrable |
| --- | --- | --- |
| Étude | 2 mois | Dossier d'exécution |
| Travaux | 18 mois | Parc rénové |
"""


@pytest.fixture
def document(db, company) -> ApplicationDocument:
    tender = Tender(title="Salé : rénovation de l'éclairage public", organization="Commune de Salé")
    db.add(tender)
    db.flush()
    application = Application(tender_id=tender.id)
    doc = ApplicationDocument(
        document_type=DocumentType.offre_technique, title="Offre technique — Salé", current_version=2
    )
    doc.sections.extend(
        [
            ApplicationSection(key="contexte", title="Contexte et enjeux", position=0, content_md=MARKDOWN),
            ApplicationSection(key="livrables", title="Livrables", position=1, content_md="Rapport final."),
        ]
    )
    application.documents.append(doc)
    db.add(application)
    db.flush()
    return doc


def _read(content: bytes):
    return read_docx(io.BytesIO(content))


def _texts(docx, style_prefix: str) -> list[str]:
    return [p.text for p in docx.paragraphs if p.style.name.startswith(style_prefix)]


def test_the_file_opens_and_has_one_heading_per_section(db, company, document):
    content = export_docx(document, from_company(company, []), tender_title="Salé : éclairage public")

    docx = _read(content)
    assert _texts(docx, "Heading 1") == ["Contexte et enjeux", "Livrables"]
    assert _texts(docx, "Title") == ["Offre technique — Salé"]


def test_the_cover_names_the_company_the_tender_and_the_version(db, company, document):
    docx = _read(export_docx(document, from_company(company, []), tender_title="Salé : éclairage public"))
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "InnoSustain" in text and "Salé : éclairage public" in text and "Version 2" in text


def test_markdown_becomes_word_paragraphs_lists_and_tables(db, company, document):
    docx = _read(export_docx(document, from_company(company, [])))

    assert "Notre approche" in _texts(docx, "Heading 3")
    assert "Audit énergétique préalable" in _texts(docx, "List Bullet")
    assert "Première phase" in _texts(docx, "List Number")
    table = docx.tables[0]
    assert len(table.rows) == 3 and table.cell(0, 0).text == "Phase"  # l'en-tête et deux lignes
    assert table.cell(2, 2).text == "Parc rénové"
    bold = [r.text for p in docx.paragraphs for r in p.runs if r.bold]
    assert "rénovation complète" in bold


def test_a_draft_warns_and_carries_no_ai_mention(db, company, document):
    docx = _read(export_docx(document, from_company(company, [])))
    text = "\n".join(p.text for p in docx.paragraphs)
    assert DRAFT_NOTICE in text
    assert AI_NOTICE not in "\n".join(p.text for p in docx.sections[0].footer.paragraphs)


def test_a_validated_document_carries_the_ai_mention_in_the_footer(db, company, document):
    document.status = AppDocStatus.validated
    db.flush()

    docx = _read(export_docx(document, from_company(company, [])))

    footer = "\n".join(p.text for p in docx.sections[0].footer.paragraphs)
    assert footer.startswith(AI_NOTICE) and "validé le" in footer
    assert DRAFT_NOTICE not in "\n".join(p.text for p in docx.paragraphs)


def test_unknown_markup_is_kept_rather_than_lost(db, company):
    docx = read_docx()
    write_markdown(docx, "> Citation conservée\n\n`code inline`\n")
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "> Citation conservée" in text and "`code inline`" in text


def test_the_export_key_names_the_document_and_its_version(document):
    key = export_key(document)
    assert key == f"applications/{document.application_id}/{document.id}-v2.docx"
