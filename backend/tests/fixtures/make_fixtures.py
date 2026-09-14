"""Génère les fichiers d'exemple de `tests/fixtures/` (PDF, DOCX, XLSX, TXT, ZIP).

Les fichiers sont versionnés dans git ; ce script ne sert qu'à les (re)créer :

    uv run python tests/fixtures/make_fixtures.py

Contenu volontairement InnoSustain (« Attestation fiscale InnoSustain 2026 ») : les phases 6 et 8
réutilisent ces fichiers pour tester l'extraction de texte et l'indexation.
"""

import io
import zipfile
from pathlib import Path

import openpyxl
import pymupdf
from docx import Document

HERE = Path(__file__).resolve().parent
TITLE = "Attestation fiscale InnoSustain 2026"
COMPANY = "Innovative & Sustainable Solutions (InnoSustain)"


def make_pdf() -> bytes:
    """Une page A4 avec un titre et trois lignes de texte extractible."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 90), TITLE, fontsize=16)
    page.insert_text((72, 130), f"Société : {COMPANY}", fontsize=11)
    page.insert_text(
        (72, 150), "Le contribuable est en situation régulière au regard de ses obligations.", fontsize=11
    )
    page.insert_text((72, 170), "Valable jusqu'au 31/12/2026.", fontsize=11)
    data = doc.tobytes(garbage=4, deflate=True)
    doc.close()
    return data


def make_docx() -> bytes:
    d = Document()
    d.add_heading("Présentation InnoSustain", level=1)
    d.add_paragraph(
        f"{COMPANY} accompagne les organisations dans leurs projets environnement, énergie et conseil."
    )
    d.add_paragraph("Références : audit énergétique (Office National X), plan climat territorial (Ville Y).")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Références"
    ws.append(["Projet", "Client", "Secteur", "Année"])
    ws.append(["Audit énergétique", "Office National X", "Énergie", 2025])
    ws.append(["Plan climat territorial", "Ville Y", "Environnement", 2024])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_txt() -> bytes:
    return "Note interne InnoSustain\nCertification ISO 14001 valide ; ISO 9001 à renouveler.\n".encode()


def make_zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members.items():
            # Horodatage fixe : l'archive ne change pas d'un lancement à l'autre.
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            zf.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)
    return buf.getvalue()


def main() -> None:
    pdf, txt = make_pdf(), make_txt()
    files = {
        "sample.pdf": pdf,
        "sample.docx": make_docx(),
        "sample.xlsx": make_xlsx(),
        "sample.txt": txt,
        "sample.zip": make_zip({"sample.txt": txt, "sample.pdf": pdf}),
    }
    for name, data in files.items():
        (HERE / name).write_bytes(data)
        print(f"{name}: {len(data)} octets")


if __name__ == "__main__":
    main()
